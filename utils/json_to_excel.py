#!/usr/bin/env python3
"""
json2excel.py

Usage:
    python json2excel.py input.json output.xlsx --array-handling join
    python json2excel.py input.jsonl output.xlsx --array-handling explode

Dependencies:
    pip install pandas openpyxl python-dateutil
"""

import argparse
import json
import math
from dateutil.parser import parse as dateparse, ParserError
from datetime import datetime
from typing import Any, Dict, List
import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl import load_workbook


def read_json_input(path: str) -> List[Dict]:
    """Read JSON file supporting .jsonl or .json (list or single object)."""
    if path.lower().endswith(".jsonl") or _is_jsonl(path):
        records = []
        with open(path, "r", encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if not line:
                    continue
                records.append(json.loads(line))
        return records
    else:
        with open(path, "r", encoding="utf-8") as f:
            obj = json.load(f)
            if isinstance(obj, list):
                return obj
            elif isinstance(obj, dict):
                return [obj]
            else:
                raise ValueError("Unsupported JSON root type: expected list or dict.")


def _is_jsonl(path: str) -> bool:
    try:
        with open(path, "r", encoding="utf-8") as f:
            first = f.readline()
            return first.strip().startswith("{") and ("\n" in first or f.readline())
    except Exception:
        return False


def flatten_dict(d: Dict[str, Any], parent_key="", sep=".") -> Dict[str, Any]:
    """Flatten nested dicts. Lists are kept as lists (handled later)."""
    items = {}
    for k, v in d.items():
        new_key = f"{parent_key}{sep}{k}" if parent_key else k
        if isinstance(v, dict):
            items.update(flatten_dict(v, new_key, sep=sep))
        else:
            items[new_key] = v
    return items


def explode_records(
    records: List[Dict], array_keys_hint: List[str] = None
) -> List[Dict]:
    """
    Expand records so that if a field is a list, we create multiple rows,
    one per element. If multiple list-fields exist in same record, it does a
    Cartesian expansion (common expected behavior).
    """
    out = []
    for rec in records:
        # find keys that are lists
        list_keys = [k for k, v in rec.items() if isinstance(v, list) and v]
        # optional hint: if provided only explode those keys
        if array_keys_hint:
            list_keys = [k for k in list_keys if k in array_keys_hint]
        if not list_keys:
            out.append(rec)
            continue

        # Start with one base row
        rows = [rec.copy()]
        for key in list_keys:
            new_rows = []
            for row in rows:
                vals = row.get(key, [])
                # if empty list, produce one row with None
                if not vals:
                    r = row.copy()
                    r[key] = None
                    new_rows.append(r)
                else:
                    for v in vals:
                        r = row.copy()
                        r[key] = v
                        new_rows.append(r)
            rows = new_rows
        out.extend(rows)
    return out


def join_array_values(v: Any, sep="; ") -> Any:
    if isinstance(v, list):
        # convert non-string elements to JSON-friendly strings
        return sep.join([_to_cell_string(x) for x in v])
    return v


def _to_cell_string(x):
    if x is None:
        return ""
    if isinstance(x, (dict, list)):
        # compact JSON
        return json.dumps(x, ensure_ascii=False, separators=(",", ":"))
    return str(x)


def coerce_types_for_excel(df: pd.DataFrame) -> pd.DataFrame:
    """
    Try to coerce columns to numeric or datetime where appropriate.
    This improves Excel cell typing (numbers/dates instead of strings).
    """
    for col in df.columns:
        # skip purely null columns
        if df[col].isnull().all():
            continue

        # try numeric
        try:
            coerced = pd.to_numeric(df[col], errors="coerce")
            # if many values converted (not all NaN), accept numeric
            converted_pct = coerced.notna().sum() / max(1, df[col].notna().sum())
            if converted_pct > 0.6:
                df[col] = coerced
                continue
        except Exception:
            pass

        # try datetime (ISO-like / parseable)
        parse_success = []
        new_vals = []
        for v in df[col]:
            if pd.isna(v):
                new_vals.append(None)
                parse_success.append(False)
                continue
            if isinstance(v, (int, float)):
                parse_success.append(False)
                new_vals.append(v)
                continue
            if isinstance(v, datetime):
                parse_success.append(True)
                new_vals.append(v)
                continue
            s = str(v).strip()
            # quick heuristic: common datetime characters
            if any(c in s for c in ("-", "T", ":", "/")) and len(s) >= 6:
                try:
                    dt = dateparse(s)
                    new_vals.append(dt)
                    parse_success.append(True)
                    continue
                except ParserError:
                    pass
            parse_success.append(False)
            new_vals.append(v)
        if sum(parse_success) / max(1, len(parse_success)) > 0.6:
            df[col] = pd.Series(new_vals)
    return df


def write_excel(df: pd.DataFrame, out_path: str, sheet_name="Sheet1"):
    # Use pandas to_excel (openpyxl)
    df.to_excel(out_path, index=False, sheet_name=sheet_name, engine="openpyxl")

    # post-process with openpyxl: set column widths, number formats for dates
    wb = load_workbook(out_path)
    ws = wb[sheet_name]

    # set widths based on max length in column (cap width)
    for i, col in enumerate(df.columns, start=1):
        max_length = max(
            (
                len(str(cell.value)) if cell.value is not None else 0
                for cell in ws[get_column_letter(i)]
            ),
            default=0,
        )
        # cap widths between 10 and 60
        adjusted_width = min(max(10, int(max_length) + 2), 60)
        ws.column_dimensions[get_column_letter(i)].width = adjusted_width

        # try to set date format if column dtype is datetime
        try:
            col_series = df[col]
            if (
                pd.api.types.is_datetime64_any_dtype(col_series)
                or col_series.apply(lambda x: isinstance(x, datetime)).any()
            ):
                for cell in ws[get_column_letter(i)][1:]:  # skip header
                    if isinstance(cell.value, datetime):
                        cell.number_format = "yyyy-mm-dd hh:mm:ss"
        except Exception:
            pass

    wb.save(out_path)


def json_to_excel(
    in_path: str,
    out_path: str,
    array_handling: str = "join",  # join | explode
    array_join_sep: str = "; ",
    explode_array_keys: List[str] = None,
):
    raw = read_json_input(in_path)
    # flatten nested dicts
    flattened = [flatten_dict(r) for r in raw]

    if array_handling == "explode":
        # first, ensure lists are kept as lists in flattened dicts
        exploded = explode_records(flattened, array_keys_hint=explode_array_keys)
        processed = []
        for r in exploded:
            # for any remaining lists, convert to string (join)
            processed.append(
                {
                    k: (
                        join_array_values(v, array_join_sep)
                        if isinstance(v, list)
                        else v
                    )
                    for k, v in r.items()
                }
            )
    else:
        # join arrays into string cells
        processed = []
        for r in flattened:
            processed.append(
                {
                    k: (
                        join_array_values(v, array_join_sep)
                        if isinstance(v, list)
                        else v
                    )
                    for k, v in r.items()
                }
            )

    df = pd.DataFrame(processed)
    df = coerce_types_for_excel(df)
    write_excel(df, out_path)


def main():
    parser = argparse.ArgumentParser(
        description="Convert JSON/JSONL to Excel without format misalignment."
    )
    parser.add_argument("input", help="Input JSON or JSONL file path")
    parser.add_argument("output", help="Output .xlsx file path")
    parser.add_argument(
        "--array-handling",
        choices=["join", "explode"],
        default="join",
        help="How to handle arrays: join (single cell) or explode (one row per array element). Default join.",
    )
    parser.add_argument(
        "--array-join-sep",
        default="; ",
        help="Separator to join array elements when using join (default '; ').",
    )
    parser.add_argument(
        "--explode-keys",
        nargs="*",
        default=None,
        help="Optional list of specific dot-keys to explode (only with explode mode).",
    )
    args = parser.parse_args()

    json_to_excel(
        args.input,
        args.output,
        array_handling=args.array_handling,
        array_join_sep=args.array_join_sep,
        explode_array_keys=args.explode_keys,
    )
    print(f"Wrote Excel file: {args.output}")


if __name__ == "__main__":
    main()
